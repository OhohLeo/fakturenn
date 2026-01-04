"""Tests for XLS source implementation."""

import os
import tempfile
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from app.sources.xls import XLSSource


@pytest.fixture
def sample_xlsx_file():
    """Create a sample XLSX file for testing."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # Header row
    ws["A1"] = "Invoice ID"
    ws["B1"] = "Date"
    ws["C1"] = "Amount"
    ws["D1"] = "Description"

    # Data rows
    ws["A2"] = "INV-001"
    ws["B2"] = date(2024, 1, 15)
    ws["C2"] = "100,00 €"
    ws["D2"] = "Test invoice 1"

    ws["A3"] = "INV-002"
    ws["B3"] = date(2024, 2, 20)
    ws["C3"] = "250,50 €"
    ws["D3"] = "Test invoice 2"

    ws["A4"] = "INV-003"
    ws["B4"] = date(2024, 3, 10)
    ws["C4"] = "75,00 €"
    ws["D4"] = "Test invoice 3"

    # Save to temp file
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        wb.save(f.name)
        yield f.name

    # Cleanup
    os.unlink(f.name)


@pytest.fixture
def xls_source(sample_xlsx_file):
    """Create XLS source with sample file."""
    config = {
        "file_path": sample_xlsx_file,
        "sheet_name": "Invoices",
        "skip_rows": 1,  # Skip header
    }
    return XLSSource(config)


@pytest.mark.asyncio
async def test_xls_source_connect(xls_source):
    """Test XLS source connection."""
    await xls_source.connect()
    assert xls_source.is_connected
    await xls_source.disconnect()


@pytest.mark.asyncio
async def test_xls_source_list_items(xls_source):
    """Test listing items from XLS source."""
    await xls_source.connect()

    rows = await xls_source.list_items()

    assert len(rows) == 3
    assert rows[0].row_number == 2  # First data row (after header)
    assert rows[0].data["A"] == "INV-001"

    await xls_source.disconnect()


@pytest.mark.asyncio
async def test_xls_source_list_items_with_limit(xls_source):
    """Test listing items with limit."""
    await xls_source.connect()

    rows = await xls_source.list_items(limit=2)

    assert len(rows) == 2

    await xls_source.disconnect()


@pytest.mark.asyncio
async def test_xls_source_get_item(xls_source):
    """Test getting a single item by row number."""
    await xls_source.connect()

    row = await xls_source.get_item("3")  # Row 3

    assert row.row_number == 3
    assert row.data["A"] == "INV-002"

    await xls_source.disconnect()


@pytest.mark.asyncio
async def test_xls_source_date_filter(sample_xlsx_file):
    """Test filtering rows by date."""
    config = {
        "file_path": sample_xlsx_file,
        "sheet_name": "Invoices",
        "skip_rows": 1,
    }
    source = XLSSource(config)
    await source.connect()

    # Filter rows after 2024-02-01
    filters = {
        "date_after": "2024-02-01",
        "date_column": "B",
    }
    rows = await source.list_items(filters=filters)

    # Should get INV-002 and INV-003 (dates in Feb and March)
    assert len(rows) == 2

    await source.disconnect()


@pytest.mark.asyncio
async def test_xls_source_column_filter(sample_xlsx_file):
    """Test filtering rows by column value."""
    config = {
        "file_path": sample_xlsx_file,
        "sheet_name": "Invoices",
        "skip_rows": 1,
    }
    source = XLSSource(config)
    await source.connect()

    # Filter by specific invoice ID
    filters = {
        "column_equals": {"A": "INV-002"},
    }
    rows = await source.list_items(filters=filters)

    assert len(rows) == 1
    assert rows[0].data["A"] == "INV-002"

    await source.disconnect()


@pytest.mark.asyncio
async def test_xls_source_file_not_found():
    """Test error handling for missing file."""
    config = {
        "file_path": "/nonexistent/path/file.xlsx",
        "sheet_name": "Sheet1",
    }
    source = XLSSource(config)

    with pytest.raises(Exception):
        await source.connect()


@pytest.mark.asyncio
async def test_xls_source_context_manager(sample_xlsx_file):
    """Test using XLS source as context manager."""
    config = {
        "file_path": sample_xlsx_file,
        "sheet_name": "Invoices",
        "skip_rows": 1,
    }

    async with XLSSource(config) as source:
        rows = await source.list_items()
        assert len(rows) == 3

    # Should be disconnected after context
    assert not source.is_connected
