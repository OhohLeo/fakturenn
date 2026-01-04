"""Tests for XLS2Record parser implementation."""

import os
import tempfile
from datetime import date
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import Workbook

from app.parsers.xls2record import XLS2RecordParser
from app.sources.models import RawRow
from app.sources.xls import XLSSource


@pytest.fixture
def sample_xlsx_file():
    """Create a sample XLSX file for testing."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # Header row
    ws["A1"] = "Invoice ID"
    ws["B1"] = "Date"
    ws["C1"] = "Amount"
    ws["D1"] = "Description"

    # Data rows
    ws["A2"] = "INV-001"
    ws["B2"] = date(2024, 1, 15)
    ws["C2"] = "100,00 €"
    ws["D2"] = "Test invoice 1"

    ws["A3"] = "INV-002"
    ws["B3"] = date(2024, 2, 20)
    ws["C3"] = "250,50 €"
    ws["D3"] = "Test invoice 2"

    # Save to temp file
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        wb.save(f.name)
        yield f.name

    # Cleanup
    os.unlink(f.name)


@pytest.fixture
def mock_xls_source():
    """Create a mock XLS source."""
    source = MagicMock(spec=XLSSource)
    source.list_items = AsyncMock(
        return_value=[
            RawRow(
                row_number=2,
                data={
                    "A": "INV-001",
                    "B": date(2024, 1, 15),
                    "C": "100,00 €",
                    "D": "Test invoice 1",
                },
                sheet_name="Invoices",
                file_path="/test/file.xlsx",
            ),
            RawRow(
                row_number=3,
                data={
                    "A": "INV-002",
                    "B": date(2024, 2, 20),
                    "C": "250,50 €",
                    "D": "Test invoice 2",
                },
                sheet_name="Invoices",
                file_path="/test/file.xlsx",
            ),
        ]
    )
    return source


@pytest.mark.asyncio
async def test_xls2record_parser_basic(mock_xls_source):
    """Test basic XLS to record parsing."""
    rules = {
        "column_mapping": {
            "A": "invoice_id",
            "B": "date",
            "C": "amount_text",
            "D": "description",
        },
        "date_column": "B",
    }

    parser = XLS2RecordParser(source=mock_xls_source, rules=rules)
    records = await parser.parse()

    assert len(records) == 2
    assert records[0].data["invoice_id"] == "INV-001"
    assert records[0].data["amount_text"] == "100,00 €"
    assert records[0].record_date == date(2024, 1, 15)


@pytest.mark.asyncio
async def test_xls2record_parser_with_filters(mock_xls_source):
    """Test parsing with filters."""
    rules = {
        "column_mapping": {
            "A": "invoice_id",
            "B": "date",
            "C": "amount_text",
        },
        "date_column": "B",
        "filters": {
            "date_after": "2024-02-01",
        },
    }

    parser = XLS2RecordParser(source=mock_xls_source, rules=rules)
    await parser.parse()

    # Should have called list_items with filters
    mock_xls_source.list_items.assert_called_once()
    call_kwargs = mock_xls_source.list_items.call_args
    assert "filters" in call_kwargs.kwargs


@pytest.mark.asyncio
async def test_xls2record_parser_skip_empty_rows():
    """Test skipping empty rows."""
    source = MagicMock(spec=XLSSource)
    source.list_items = AsyncMock(
        return_value=[
            RawRow(
                row_number=2,
                data={"A": "", "B": None, "C": "", "D": ""},
                sheet_name="Sheet1",
                file_path="/test.xlsx",
            ),
            RawRow(
                row_number=3,
                data={
                    "A": "INV-001",
                    "B": date(2024, 1, 15),
                    "C": "100,00",
                    "D": "Test",
                },
                sheet_name="Sheet1",
                file_path="/test.xlsx",
            ),
        ]
    )

    rules = {
        "column_mapping": {"A": "invoice_id", "B": "date", "C": "amount", "D": "desc"},
        "date_column": "B",
        "skip_empty_rows": True,
    }

    parser = XLS2RecordParser(source=source, rules=rules)
    records = await parser.parse()

    # Should skip empty row, only return one record
    assert len(records) == 1
    assert records[0].data["invoice_id"] == "INV-001"


@pytest.mark.asyncio
async def test_xls2record_parser_trim_whitespace():
    """Test whitespace trimming."""
    source = MagicMock(spec=XLSSource)
    source.list_items = AsyncMock(
        return_value=[
            RawRow(
                row_number=2,
                data={
                    "A": "  INV-001  ",
                    "B": date(2024, 1, 15),
                    "C": " 100,00 ",
                    "D": "  Test  ",
                },
                sheet_name="Sheet1",
                file_path="/test.xlsx",
            ),
        ]
    )

    rules = {
        "column_mapping": {"A": "invoice_id", "B": "date", "C": "amount", "D": "desc"},
        "date_column": "B",
        "trim_whitespace": True,
    }

    parser = XLS2RecordParser(source=source, rules=rules)
    records = await parser.parse()

    assert records[0].data["invoice_id"] == "INV-001"
    assert records[0].data["amount"] == "100,00"
    assert records[0].data["desc"] == "Test"


@pytest.mark.asyncio
async def test_xls2record_parser_date_formats():
    """Test parsing various date formats."""
    source = MagicMock(spec=XLSSource)
    source.list_items = AsyncMock(
        return_value=[
            RawRow(
                row_number=2,
                data={"A": "INV-001", "B": "15/01/2024", "C": "100"},
                sheet_name="Sheet1",
                file_path="/test.xlsx",
            ),
        ]
    )

    rules = {
        "column_mapping": {"A": "invoice_id", "B": "date", "C": "amount"},
        "date_column": "B",
        "date_format": "%d/%m/%Y",
    }

    parser = XLS2RecordParser(source=source, rules=rules)
    records = await parser.parse()

    assert records[0].record_date == date(2024, 1, 15)


@pytest.mark.asyncio
async def test_xls2record_parser_no_column_mapping():
    """Test error when no column mapping defined."""
    source = MagicMock(spec=XLSSource)
    source.list_items = AsyncMock(
        return_value=[
            RawRow(
                row_number=2,
                data={"A": "INV-001", "B": date(2024, 1, 15)},
                sheet_name="Sheet1",
                file_path="/test.xlsx",
            ),
        ]
    )

    rules = {}  # No column_mapping

    parser = XLS2RecordParser(source=source, rules=rules)
    records = await parser.parse()

    # Should skip rows that fail extraction
    assert len(records) == 0


@pytest.mark.asyncio
async def test_xls2record_parser_metadata(mock_xls_source):
    """Test that extraction metadata is included."""
    rules = {
        "column_mapping": {"A": "invoice_id", "B": "date"},
        "date_column": "B",
    }

    parser = XLS2RecordParser(source=mock_xls_source, rules=rules)
    records = await parser.parse()

    assert records[0].extraction_metadata["extraction_method"] == "column_mapping"
    assert records[0].extraction_metadata["file_path"] == "/test/file.xlsx"
    assert records[0].extraction_metadata["sheet_name"] == "Invoices"
